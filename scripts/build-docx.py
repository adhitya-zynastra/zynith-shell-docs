#!/usr/bin/env python3
"""Compile the canonical Markdown into ZynithShell-Technical-Report.docx.

Cover page, metadata, table of contents field, page numbers in the footer, figure and table captions,
monospaced code blocks, real Word tables. Content comes only from the Markdown sources.
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docbuild_common as C

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x1b, 0x1b, 0x1f)
MUTE = RGBColor(0x5b, 0x60, 0x70)
ACCENT = RGBColor(0x30, 0x4f, 0xd6)

def field(par, instr):
    r = par.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin"); r._r.append(fld)
    r2 = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = instr; r2._r.append(it)
    r3 = par.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate"); r3._r.append(f2)
    r4 = par.add_run("…")
    r5 = par.add_run(); f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end"); r5._r.append(f3)

def footer_pagenum(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(p, "PAGE")
    for r in p.runs:
        r.font.size = Pt(8); r.font.color.rgb = MUTE

def style(doc):
    n = doc.styles["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(10.5); n.font.color.rgb = INK
    n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.12
    for i, size in ((1, 20), (2, 15), (3, 12.5), (4, 11)):
        s = doc.styles[f"Heading {i}"]
        s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True
        s.font.color.rgb = INK if i > 1 else ACCENT
        s.paragraph_format.space_before = Pt(14 if i < 3 else 10); s.paragraph_format.space_after = Pt(4)

def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Consolas"; r.font.size = Pt(8.4)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    return p

def add_table(doc, rows, caption_n):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols); t.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(cols):
            txt = C.plain(row[ci]) if ci < len(row) else ""
            cells[ci].text = txt
            for par in cells[ci].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8.6)
                    if ri == 0: run.font.bold = True
    cap = doc.add_paragraph(f"Table {caption_n}")
    cap.runs[0].font.size = Pt(8); cap.runs[0].font.color.rgb = MUTE
    cap.paragraph_format.space_after = Pt(10)

def rich(par, text):
    """Bold/code/italic inline spans."""
    import re
    for tok in re.split(C.TOKEN.pattern, text):
        if not tok: continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(C.plain(tok)); r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok.strip("`")); r.font.name = "Consolas"; r.font.size = Pt(9.2)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(C.plain(tok)); r.italic = True
        else:
            par.add_run(C.plain(tok))

def main():
    meta = C.meta()
    doc = Document()
    style(doc)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    sec.left_margin = sec.right_margin = Inches(0.95)
    footer_pagenum(sec)

    # ---- cover ----
    for _ in range(4): doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("ZYNITH SHELL"); r.font.size = Pt(38); r.bold = True; r.font.color.rgb = ACCENT
    s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("An Engineering Record"); r.font.size = Pt(14); r.font.color.rgb = INK
    s2 = doc.add_paragraph(); s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s2.add_run("Fedora 44 · niri 26.04 · patched Noctalia 5.1.0"); r.font.size = Pt(11); r.font.color.rgb = MUTE
    for _ in range(2): doc.add_paragraph()
    info = [("Documentation version", meta["documentation_version"]),
            ("Generated", meta["generated"]),
            ("Source commit", meta["source_commit"]),
            ("Last audited phase", str(meta["last_audited_phase"])),
            ("Machine", meta["machine"]["cpu"] + " · " + meta["machine"]["gpu"]),
            ("Display", meta["machine"]["display"]),
            ("Author", "M.S.Adhitya"),
            ("Engineering assistant", "Claude Code (Anthropic)")]
    tb = doc.add_table(rows=0, cols=2); tb.style = "Light List Accent 1"
    for k, v in info:
        c = tb.add_row().cells
        c[0].text = k; c[1].text = v
        for cell in c:
            for par in cell.paragraphs:
                for run in par.runs: run.font.size = Pt(9.5)
    doc.add_paragraph()
    fw = doc.add_paragraph()
    r = fw.add_run(C.FOREWORD); r.font.size = Pt(9); r.font.color.rgb = MUTE
    doc.add_page_break()

    # ---- table of contents ----
    h = doc.add_paragraph("Contents"); h.style = doc.styles["Heading 1"]
    p = doc.add_paragraph(); field(p, r'TOC \o "1-3" \h \z \u')
    doc.add_paragraph().add_run(
        "This table of contents is a Word field: open the document and press F9 (or Ctrl+A, F9) to populate it."
    ).font.size = Pt(8.5)
    doc.add_page_break()

    fig_n, tab_n = 0, 0
    for section, files in C.REPORT:
        doc.add_heading(section, 1)
        for path in files:
            if not os.path.exists(os.path.join(C.ROOT, path)):
                continue
            first = True
            for b in C.blocks(path):
                kind = b[0]
                if kind == "h":
                    lvl = min(b[1] + 1, 4)
                    if first and b[1] == 1:
                        doc.add_heading(C.plain(b[2]), 2); first = False
                    else:
                        doc.add_heading(C.plain(b[2]), lvl)
                elif kind == "p":
                    par = doc.add_paragraph(); rich(par, b[1])
                elif kind == "li":
                    par = doc.add_paragraph(style="List Bullet"); rich(par, b[1])
                elif kind == "oli":
                    par = doc.add_paragraph(style="List Number"); rich(par, b[2])
                elif kind == "quote":
                    par = doc.add_paragraph(); par.paragraph_format.left_indent = Inches(0.3)
                    rich(par, b[1])
                    for run in par.runs: run.italic = True; run.font.color.rgb = MUTE
                elif kind == "code":
                    code(doc, b[1])
                elif kind == "table":
                    tab_n += 1; add_table(doc, b[1], tab_n)
            doc.add_paragraph()
        # figures after the architecture and performance sections
        if section in ("Architecture", "Performance Engineering"):
            wanted = C.FIGURES[:5] if section == "Architecture" else C.FIGURES[5:]
            for rel, cap in wanted:
                png = rel.replace(".svg", ".png")
                full = os.path.join(C.ROOT, png)
                if not os.path.exists(full): continue
                fig_n += 1
                doc.add_picture(full, width=Inches(6.1))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp = doc.add_paragraph(f"Figure {fig_n}. {cap}")
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.runs[0].font.size = Pt(8.5); cp.runs[0].font.color.rgb = MUTE
        doc.add_page_break()

    # ---- screenshots appendix ----
    doc.add_heading("Appendix — Interface", 1)
    shots = C.SHOTS
    for fn, cap in shots:
        full = os.path.join(C.ROOT, "07_Assets", "screenshots", fn)
        if not os.path.exists(full): continue
        fig_n += 1
        doc.add_picture(full, width=Inches(5.9))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp = doc.add_paragraph(f"Figure {fig_n}. {cap}")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.runs[0].font.size = Pt(8.5); cp.runs[0].font.color.rgb = MUTE

    out = os.path.join(C.ROOT, "99_Reports", "generated", "ZynithShell-Technical-Report.docx")
    doc.save(out)
    print("DOCX written:", out, f"({fig_n} figures, {tab_n} tables)")

if __name__ == "__main__":
    main()
